import streamlit as st
import pandas as pd
import io
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment

# ==========================================
# ⚙️ ตั้งค่าหน้าเว็บ
# ==========================================
st.set_page_config(page_title="โปรแกรมสรุปยอดรายได้และกำไร", page_icon="💰", layout="wide")
st.title("📊 โปรแกรมสรุปยอดรายได้ และ คำนวณกำไร (Reconciliation)")

# ==========================================
# 📂 ส่วนที่ 1: อัปโหลดไฟล์ต้นทุน
# ==========================================
st.markdown("### 📂 1. อัปโหลดไฟล์ต้นทุนสินค้า (Real Cost)")
st.info("💡 กรุณาอัปโหลดไฟล์ต้นทุนก่อนเป็นอันดับแรก เพื่อให้ระบบดึงต้นทุนไปคำนวณกำไรในแต่ละออเดอร์")
cost_file = st.file_uploader("อัปโหลดไฟล์ ต้นทุน (Real Cost.xlsx) ที่นี่", type=['xlsx'])

if cost_file is not None:
    try:
        df_cost = pd.read_excel(cost_file)
        st.success("✅ โหลดไฟล์ต้นทุนสำเร็จ! ระบบพร้อมคำนวณกำไรให้ทุกแพลตฟอร์มแล้วครับ")
    except Exception as e:
        st.error(f"❌ อ่านไฟล์ต้นทุนไม่ได้ โปรดตรวจสอบไฟล์อีกครั้ง: {e}")
        df_cost = None
else:
    df_cost = None
    st.warning("⚠️ หากยังไม่อัปโหลดไฟล์ต้นทุน ระบบจะแสดงผล 'ต้นทุนรวม = 0' และกำไรจะไม่สะท้อนความเป็นจริงครับ")
    
st.markdown("---")

# ==========================================
# 💾 ฟังก์ชันช่วยเหลือ (อ่านไฟล์ + สร้าง Excel)
# ==========================================
def read_uploaded_file(uploaded_file, skiprows=None):
    if uploaded_file.name.endswith('.csv'):
        return pd.read_csv(uploaded_file)
    else:
        if skiprows:
            return pd.read_excel(uploaded_file, skiprows=skiprows)
        return pd.read_excel(uploaded_file)

def generate_excel(df, platform_name, header_color_hex):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{platform_name}_Summary"

    ws.append(list(df.columns))
    for r in dataframe_to_rows(df, index=False, header=False):
        ws.append(r)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=header_color_hex, end_color=header_color_hex, fill_type="solid")
    for col in ws.columns:
        cell = col[0]
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def display_summary_dashboard(df, order_col, income_col, cost_col):
    st.markdown("### 📈 สรุปผลประกอบการ")
    total_income = df.drop_duplicates(subset=[order_col])[income_col].sum()
    total_cost = df[cost_col].sum()
    total_profit = total_income - total_cost

    col1, col2, col3 = st.columns(3)
    col1.metric("💰 เงินเข้าสุทธิรวม", f"{total_income:,.2f} ฿")
    col2.metric("📦 ต้นทุนรวม", f"{total_cost:,.2f} ฿")
    col3.metric("✨ กำไรสุทธิ", f"{total_profit:,.2f} ฿")
    
    st.markdown("#### 🔍 ตัวอย่างข้อมูล (5 รายการแรก)")
    st.dataframe(df.head(5), use_container_width=True)

# ==========================================
# 🟠 ฟังก์ชันประมวลผล Shopee
# ==========================================
def process_shopee(file_income, file_orders, df_cost_data):
    df_income = read_uploaded_file(file_income, skiprows=17)
    df_orders = read_uploaded_file(file_orders)

    df_income = df_income.dropna(subset=['รหัสคำสั่งซื้อ']).copy()
    df_income['รหัสคำสั่งซื้อ'] = df_income['รหัสคำสั่งซื้อ'].astype(str)
    df_income['จำนวนเงิน'] = pd.to_numeric(df_income['จำนวนเงิน'], errors='coerce').fillna(0)
    net_income_df = df_income.groupby('รหัสคำสั่งซื้อ')['จำนวนเงิน'].sum().reset_index()

    df_orders['หมายเลขคำสั่งซื้อ'] = df_orders['หมายเลขคำสั่งซื้อ'].astype(str)
    
    if df_cost_data is not None and 'เลขอ้างอิง SKU (SKU Reference No.)' in df_orders.columns:
        df_orders = df_orders.merge(df_cost_data[['เลข SKU', 'ราคาต้นทุน']], left_on='เลขอ้างอิง SKU (SKU Reference No.)', right_on='เลข SKU', how='left')
        df_orders['ราคาต้นทุนต่อชิ้น'] = pd.to_numeric(df_orders['ราคาต้นทุน'], errors='coerce').fillna(0)
        df_orders['รวมต้นทุน (บาท)'] = df_orders['ราคาต้นทุนต่อชิ้น'] * pd.to_numeric(df_orders['จำนวน'], errors='coerce').fillna(0)
        df_orders.drop(columns=['เลข SKU', 'ราคาต้นทุน'], inplace=True)
    else:
        df_orders['ราคาต้นทุนต่อชิ้น'] = 0
        df_orders['รวมต้นทุน (บาท)'] = 0

    df_orders = df_orders.merge(net_income_df, left_on='หมายเลขคำสั่งซื้อ', right_on='รหัสคำสั่งซื้อ', how='left')
    df_orders['ยอดเงินเข้าจาก Shopee (บาท)'] = df_orders['จำนวนเงิน'].fillna(0)
    
    if 'รหัสคำสั่งซื้อ' in df_orders.columns: df_orders.drop(columns=['รหัสคำสั่งซื้อ'], inplace=True)
    if 'จำนวนเงิน' in df_orders.columns: df_orders.drop(columns=['จำนวนเงิน'], inplace=True)

    cols = list(df_orders.columns)
    new_cols = ['ราคาต้นทุนต่อชิ้น', 'รวมต้นทุน (บาท)', 'ยอดเงินเข้าจาก Shopee (บาท)']
    cols = [c for c in cols if c not in new_cols]
    idx = cols.index('จำนวน') + 1 if 'จำนวน' in cols else len(cols)
    cols = cols[:idx] + new_cols + cols[idx:]
    
    df_orders = df_orders[cols]
    return generate_excel(df_orders, "Shopee", "F94D29"), df_orders

# ==========================================
# ⚫ ฟังก์ชันประมวลผล TikTok
# ==========================================
def process_tiktok(file_income, file_orders, df_cost_data):
    df_income = read_uploaded_file(file_income)
    df_orders = read_uploaded_file(file_orders)

    df_income = df_income.dropna(subset=['หมายเลขคำสั่งซื้อ/การปรับ']).copy()
    df_income['หมายเลขคำสั่งซื้อ/การปรับ'] = df_income['หมายเลขคำสั่งซื้อ/การปรับ'].astype(str).str.replace(r'\.0$', '', regex=True)

    # 💡 จุดที่แก้ไข: เช็คชื่อคอลัมน์เงินเข้าว่า TikTok ใช้ชื่อไหน
    income_col_name = None
    if 'ยอดการชำระเงินทั้งหมด' in df_income.columns:
        income_col_name = 'ยอดการชำระเงินทั้งหมด'
    elif 'จำนวนเงินที่ชำระทั้งหมด' in df_income.columns:
        income_col_name = 'จำนวนเงินที่ชำระทั้งหมด'
    
    if income_col_name:
        df_income[income_col_name] = pd.to_numeric(df_income[income_col_name], errors='coerce').fillna(0)
        net_income_df = df_income.groupby('หมายเลขคำสั่งซื้อ/การปรับ')[income_col_name].sum().reset_index()
    else:
        raise ValueError("ไม่พบคอลัมน์ยอดเงินเข้าในไฟล์ TikTok (ตรวจสอบไฟล์อีกครั้ง)")

    df_orders = df_orders[pd.to_numeric(df_orders['Order ID'], errors='coerce').notnull()].copy()
    df_orders['Order ID'] = df_orders['Order ID'].astype(str).str.replace(r'\.0$', '', regex=True)
    df_orders['Quantity'] = pd.to_numeric(df_orders['Quantity'], errors='coerce').fillna(0)

    if df_cost_data is not None and 'Seller SKU' in df_orders.columns:
        df_orders = df_orders.merge(df_cost_data[['เลข SKU', 'ราคาต้นทุน']], left_on='Seller SKU', right_on='เลข SKU', how='left')
        df_orders['ราคาต้นทุนต่อชิ้น'] = pd.to_numeric(df_orders['ราคาต้นทุน'], errors='coerce').fillna(0)
        df_orders['รวมต้นทุน (บาท)'] = df_orders['ราคาต้นทุนต่อชิ้น'] * df_orders['Quantity']
        df_orders.drop(columns=['เลข SKU', 'ราคาต้นทุน'], inplace=True)
    else:
        df_orders['ราคาต้นทุนต่อชิ้น'] = 0
        df_orders['รวมต้นทุน (บาท)'] = 0

    df_orders = df_orders.merge(net_income_df, left_on='Order ID', right_on='หมายเลขคำสั่งซื้อ/การปรับ', how='left')
    df_orders['ยอดเงินเข้าจาก TikTok (บาท)'] = df_orders[income_col_name].fillna(0)
    
    if 'หมายเลขคำสั่งซื้อ/การปรับ' in df_orders.columns: df_orders.drop(columns=['หมายเลขคำสั่งซื้อ/การปรับ'], inplace=True)
    if income_col_name in df_orders.columns: df_orders.drop(columns=[income_col_name], inplace=True)

    cols = list(df_orders.columns)
    new_cols = ['ราคาต้นทุนต่อชิ้น', 'รวมต้นทุน (บาท)', 'ยอดเงินเข้าจาก TikTok (บาท)']
    cols = [c for c in cols if c not in new_cols]
    idx = cols.index('Quantity') + 1 if 'Quantity' in cols else len(cols)
    cols = cols[:idx] + new_cols + cols[idx:]
    
    df_orders = df_orders[cols]
    return generate_excel(df_orders, "TikTok", "000000"), df_orders

# ==========================================
# 🔵 ฟังก์ชันประมวลผล Lazada
# ==========================================
def process_lazada(file_income, file_orders, df_cost_data):
    df_income = read_uploaded_file(file_income)
    df_orders = read_uploaded_file(file_orders)

    df_income = df_income.dropna(subset=['หมายเลขคำสั่งซื้อ']).copy()
    df_income['หมายเลขคำสั่งซื้อ'] = df_income['หมายเลขคำสั่งซื้อ'].astype(str).str.replace(r'\.0$', '', regex=True)
    net_income_df = df_income.groupby('หมายเลขคำสั่งซื้อ')['จำนวนเงิน(รวมภาษี)'].sum().reset_index()

    df_orders['orderNumber'] = df_orders['orderNumber'].astype(str).str.replace(r'\.0$', '', regex=True)
    
    if df_cost_data is not None and 'sellerSku' in df_orders.columns:
        df_orders = df_orders.merge(df_cost_data[['เลข SKU', 'ราคาต้นทุน']], left_on='sellerSku', right_on='เลข SKU', how='left')
        df_orders['ราคาต้นทุนต่อชิ้น'] = pd.to_numeric(df_orders['ราคาต้นทุน'], errors='coerce').fillna(0)
        df_orders['รวมต้นทุน (บาท)'] = df_orders['ราคาต้นทุนต่อชิ้น'] * 1 
        df_orders.drop(columns=['เลข SKU', 'ราคาต้นทุน'], inplace=True)
    else:
        df_orders['ราคาต้นทุนต่อชิ้น'] = 0
        df_orders['รวมต้นทุน (บาท)'] = 0

    df_orders = df_orders.merge(net_income_df, left_on='orderNumber', right_on='หมายเลขคำสั่งซื้อ', how='left')
    df_orders['ยอดเงินเข้าจาก Lazada (บาท)'] = df_orders['จำนวนเงิน(รวมภาษี)'].fillna(0)
    
    if 'หมายเลขคำสั่งซื้อ' in df_orders.columns: df_orders.drop(columns=['หมายเลขคำสั่งซื้อ'], inplace=True)
    if 'จำนวนเงิน(รวมภาษี)' in df_orders.columns: df_orders.drop(columns=['จำนวนเงิน(รวมภาษี)'], inplace=True)

    cols = list(df_orders.columns)
    new_cols = ['ราคาต้นทุนต่อชิ้น', 'รวมต้นทุน (บาท)', 'ยอดเงินเข้าจาก Lazada (บาท)']
    cols = [c for c in cols if c not in new_cols]
    idx = cols.index('unitPrice') + 1 if 'unitPrice' in cols else len(cols)
    cols = cols[:idx] + new_cols + cols[idx:]
    
    df_orders = df_orders[cols]
    return generate_excel(df_orders, "Lazada", "1F4E78"), df_orders

# ==========================================
# 🖥️ หน้าจอแสดงผลแอปพลิเคชัน
# ==========================================
tab1, tab2, tab3 = st.tabs(["🟠 ยอด Shopee", "⚫ ยอด TikTok", "🔵 ยอด Lazada"])

# --- แท็บ Shopee ---
with tab1:
    st.header("🟠 สรุปยอด Shopee")
    shopee_income = st.file_uploader("1. อัปโหลดไฟล์ เงินเข้า (Revenue) Shopee", type=['xlsx', 'csv'], key="sh_inc")
    shopee_orders = st.file_uploader("2. อัปโหลดไฟล์ คำสั่งซื้อ Shopee", type=['xlsx', 'csv'], key="sh_ord")

    if shopee_income and shopee_orders:
        if st.button("🚀 ประมวลผล Shopee", type="primary"):
            with st.spinner('กำลังคำนวณข้อมูล Shopee...'):
                try:
                    result_file, df_result = process_shopee(shopee_income, shopee_orders, df_cost)
                    st.success("✅ คำนวณสำเร็จ!")
                    display_summary_dashboard(df_result, 'หมายเลขคำสั่งซื้อ', 'ยอดเงินเข้าจาก Shopee (บาท)', 'รวมต้นทุน (บาท)')
                    st.download_button(label="📥 ดาวน์โหลดไฟล์ Shopee (Excel)", data=result_file, file_name="Shopee_Cost_Revenue_Report.xlsx")
                except Exception as e:
                    st.error(f"❌ พบข้อผิดพลาด: {e}")

# --- แท็บ TikTok ---
with tab2:
    st.header("⚫ สรุปยอด TikTok")
    tiktok_income = st.file_uploader("1. อัปโหลดไฟล์ เงินเข้า TikTok", type=['xlsx', 'csv'], key="tk_inc")
    tiktok_orders = st.file_uploader("2. อัปโหลดไฟล์ คำสั่งซื้อ TikTok", type=['xlsx', 'csv'], key="tk_ord")

    if tiktok_income and tiktok_orders:
        if st.button("🚀 ประมวลผล TikTok", type="primary"):
            with st.spinner('กำลังคำนวณข้อมูล TikTok...'):
                try:
                    result_file, df_result = process_tiktok(tiktok_income, tiktok_orders, df_cost)
                    st.success("✅ คำนวณสำเร็จ!")
                    display_summary_dashboard(df_result, 'Order ID', 'ยอดเงินเข้าจาก TikTok (บาท)', 'รวมต้นทุน (บาท)')
                    st.download_button(label="📥 ดาวน์โหลดไฟล์ TikTok (Excel)", data=result_file, file_name="TikTok_Cost_Revenue_Report.xlsx")
                except Exception as e:
                    st.error(f"❌ พบข้อผิดพลาด: {e}")

# --- แท็บ Lazada ---
with tab3:
    st.header("🔵 สรุปยอด Lazada")
    lazada_income = st.file_uploader("1. อัปโหลดไฟล์ เงินเข้า Lazada", type=['xlsx', 'csv'], key="lz_inc")
    lazada_orders = st.file_uploader("2. อัปโหลดไฟล์ คำสั่งซื้อ Lazada", type=['xlsx', 'csv'], key="lz_ord")

    if lazada_income and lazada_orders:
        if st.button("🚀 ประมวลผล Lazada", type="primary"):
            with st.spinner('กำลังคำนวณข้อมูล Lazada...'):
                try:
                    result_file, df_result = process_lazada(lazada_income, lazada_orders, df_cost)
                    st.success("✅ คำนวณสำเร็จ!")
                    display_summary_dashboard(df_result, 'orderNumber', 'ยอดเงินเข้าจาก Lazada (บาท)', 'รวมต้นทุน (บาท)')
                    st.download_button(label="📥 ดาวน์โหลดไฟล์ Lazada (Excel)", data=result_file, file_name="Lazada_Cost_Revenue_Report.xlsx")
                except Exception as e:
                    st.error(f"❌ พบข้อผิดพลาด: {e}")
